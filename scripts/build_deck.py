#!/usr/bin/env python3
"""Builds deck/Mayday.pptx: the 10-slide Mayday pitch, black and honey.

Everything is drawn with native shapes (hexagons are freeform paths), so the
deck needs no image files. The one optional image is the radar screenshot at
docs/screenshots/radar.png, used on slide 5 when it exists.

    python3 scripts/build_deck.py                  # build and verify
    python3 scripts/build_deck.py --no-animations  # transitions only

Needs python-pptx (pip install python-pptx==1.0.2).
"""

import math
import os
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "deck" / "Mayday.pptx"
# The env override exists so the screenshot layout can be tested without the file.
RADAR = Path(os.environ.get("MAYDAY_RADAR", ROOT / "docs" / "screenshots" / "radar.png"))
SITE = "https://mayday-alpha-eight.vercel.app"
REPO = "https://github.com/vnmoorthy/mayday"
ANIMATE = "--no-animations" not in sys.argv

# Slide geometry, in inches. M is the page margin.
W, H, M = 13.333, 7.5, 0.75
CW = W - 2 * M
SQ3 = math.sqrt(3)

# Black and honey. Colour is for data: red maydays, honey fixes, blue rescues.
BLACK, WHITE, INK = "000000", "FFFFFF", "F4F4F2"
MUTE, DIM, LINE, LINE2 = "8F8F8A", "55554F", "343434", "6E6E68"
CELL, CELL_LINE, GHOST = "0E0E0E", "2E2E2E", "1C1C1C"
HONEY, RED, BLUE = "F5A524", "FF3B30", "58B7FF"
SANS, MONO = "Helvetica Neue", "Menlo"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

NOTES = [
    "This is Mayday: the stop signal for agents. We built it for the Supabase Select hackathon, "
    "where the prompt was to build something agents want. Watch the comb: one cell goes red, and the hive learns from it.",
    "A honeybee that is attacked at a flower flies home and gives its nestmates a stop signal: a short pulse that says, "
    "stop sending foragers down that path. One bee pays the cost. The rest of the hive doesn't.",
    "Agents have no stop signal. Thousands of them hit the same Stripe webhook error, the same Supabase row-level "
    "security wall, the same Next.js error, every single day. Every agent pays again, and the vendor never finds out.",
    "Here is the loop. An agent goes down and sends a mayday. Postgres finds the crash site and sends back a briefing: "
    "the fix that worked for earlier agents, with the vendor's official fix first. When the rescue is confirmed, the best fix rises.",
    "This is live right now, so let's watch one go down. I'll send a real error and you'll see it land on the map. "
    "The numbers include charted sites, and live and test-flight traffic is always marked.",
    "Every vendor gets a tower. They see where agents crash on their product, ranked by agents down and hours lost. "
    "They pin the official fix at the exact crash site, and they pay per rescue through Stripe, only when the fix works.",
    "Every vendor also gets an airworthiness rating, and it can't be bought. It is computed from real crashes and rescues, "
    "so it moves only when agents stop going down. Agents check it before they build, with mayday preflight.",
    "How it's built. Supabase Postgres does the matching, row-level security is the permission model, and Realtime is the UI. "
    "Vercel runs Next.js 16 and the MCP server, Stripe handles Checkout and pay-per-rescue meters, and Claude agents fly the "
    "test flights. A mayday is one SQL transaction.",
    "Vendors already pay to stop developers failing on their product. Mayday lets them find the crash sites, fix them at "
    "the moment of failure, and prove it with a rating and a per-rescue bill. And there is day-one value with zero network: "
    "launch a test flight.",
    "Every agent that goes down should be the last one to go down there. That is Mayday. It's live at this address, "
    "the code is on GitHub, and your agents can connect today. Thank you.",
]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
STEPS = {}      # slide_id -> shape ids that fade in, in order
WARNINGS = []   # text that probably overflows its box


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------- primitives

def paint(shp, fill=None, line=None, lw=0.75, dash=False):
    """Flat fill and hairline. Dropping p:style removes the theme's shadow."""
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)
    if hasattr(shp, "fill"):
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    return shp


def hexagon(s, cx, cy, r, fill=None, line=None, lw=0.75):
    """A pointy-top hexagon, the same shape as the .hex clip-path in the app."""
    pts = [
        (Inches(cx + r * math.cos(math.radians(60 * i - 90))), Inches(cy + r * math.sin(math.radians(60 * i - 90))))
        for i in range(6)
    ]
    fb = s.shapes.build_freeform(pts[0][0], pts[0][1])
    fb.add_line_segments(pts[1:], close=True)
    return paint(fb.convert_to_shape(), fill, line, lw)


def polyline(s, pts, color, lw=0.75, dash=False):
    e = [(Inches(x), Inches(y)) for x, y in pts]
    fb = s.shapes.build_freeform(e[0][0], e[0][1])
    fb.add_line_segments(e[1:], close=False)
    return paint(fb.convert_to_shape(), None, color, lw, dash)


def rect(s, x, y, w, h, fill=None, line=LINE, lw=0.75):
    return paint(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)), fill, line, lw)


def dot(s, cx, cy, r, fill):
    return paint(s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r)), fill)


def rule(s, x1, y1, x2, y2, color=LINE, lw=0.75, arrow=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    paint(c, None, color, lw)
    if arrow:
        ln = c._element.spPr.find(qn("a:ln"))
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def text(s, x, y, w, h, paras, size=16, color=INK, font=SANS, track=0.0, align="l", anchor="t", line=1.0, gap=0, shape=None):
    """A text box. paras is a string, or a list of paragraphs; a paragraph is a
    string or a list of (text, overrides) runs. track is letter spacing in em."""
    box = shape or s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.left, box.top, box.width, box.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    need = 0.0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[align]
        p.line_spacing = line
        if i and gap:
            p.space_before = Pt(gap)
        runs = [(para, {})] if isinstance(para, str) else para
        width_pt, biggest = 0.0, 0.0
        for t, o in runs:
            sz, fnt, tr = o.get("size", size), o.get("font", font), o.get("track", track)
            r = p.add_run()
            r.text = t
            r.font.size = Pt(sz)
            r.font.name = fnt
            r.font.bold = False
            r.font.color.rgb = rgb(o.get("color", color))
            if tr:
                r._r.get_or_add_rPr().set("spc", str(int(round(tr * sz * 100))))
            width_pt += len(t) * sz * ((0.602 if fnt == MONO else 0.5) + tr)
            biggest = max(biggest, sz)
        # Rough fit check: wrapped lines times line height against the box.
        lines = max(1, math.ceil(width_pt * 1.08 / (w * 72)))
        need += (lines * biggest * 1.2 * line + (gap if i else 0)) / 72
    if need > h + 0.04:
        WARNINGS.append(f"text may overflow ({need:.2f}in needed, {h:.2f}in box): {str(paras)[:60]}")
    return box


def step(s, shapes):
    """Group shapes into one unit that fades in as the next animation step."""
    g = s.shapes.add_group_shape([x for x in shapes if x is not None])
    STEPS.setdefault(s.slide_id, []).append(g.shape_id)
    return g


def comb(n):
    """Axial coordinates of a hexagonal comb of radius n, with each cell's ring."""
    for q in range(-n, n + 1):
        for r in range(max(-n, -q - n), min(n, -q + n) + 1):
            yield q, r, max(abs(q), abs(r), abs(q + r))


def cell_xy(cx, cy, R, q, r):
    return cx + R * SQ3 * (q + r / 2), cy + R * 1.5 * r


def slide(n, section):
    s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only: keeps a real title for outline view
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BLACK)
    s.notes_slide.notes_text_frame.text = NOTES[n - 1]
    # Instrument chrome: mark and name top left, slide counter top right.
    hexagon(s, M + 0.11, 0.63, 0.12, HONEY)
    text(s, M + 0.36, 0.54, 3, 0.2, "MAYDAY", size=10, color=WHITE, font=MONO, track=0.16)
    text(s, W - M - 6, 0.54, 6, 0.2, f"{n:02d} / 10  ·  {section}", size=10, color=MUTE, font=MONO, track=0.16, align="r")
    rule(s, M, 0.98, W - M, 0.98)
    rule(s, M, 6.88, W - M, 6.88)
    text(s, M, 7.0, 6, 0.18, "THE STOP SIGNAL FOR AGENTS", size=8.5, color=DIM, font=MONO, track=0.16)
    text(s, W - M - 6, 7.0, 6, 0.18, SITE.replace("https://", ""), size=8.5, color=DIM, font=MONO, track=0.04, align="r")
    return s


def title(s, paras, x, y, w, h, size=54, color=WHITE, track=-0.035, line=0.9):
    return text(s, x, y, w, h, paras, size=size, color=color, track=track, line=line, shape=s.shapes.title)


# -------------------------------------------------------------------- slides

def s01_title():
    s = slide(1, "TITLE")
    title(s, "Mayday", M, 1.95, 6.6, 1.8, size=120, track=-0.045, line=0.85)
    text(s, M, 3.95, 6.6, 0.65, "The stop signal for agents.", size=34, color=WHITE, track=-0.03)
    rule(s, M, 4.98, M + 0.6, 4.98, HONEY, 1.5)
    text(s, M, 5.14, 6.6, 0.25, "SUPABASE SELECT 2026 HACKATHON", size=10.5, color=MUTE, font=MONO, track=0.16)

    cx, cy, R = 10.0, 3.93, 0.42
    blue = {(2, -1), (1, 1), (-1, 2), (-2, 1), (-1, -1), (1, -2)}
    for q, r, ring in comb(3):
        x, y = cell_xy(cx, cy, R, q, r)
        hexagon(s, x, y, R * 0.92, CELL if ring < 3 else None, CELL_LINE if ring < 3 else GHOST)
    # One cell goes red, its neighbours turn honey, then the rescues arrive.
    step(s, [hexagon(s, cx, cy, R * 0.92, RED)])
    step(s, [hexagon(s, *cell_xy(cx, cy, R, q, r), R * 0.92, HONEY) for q, r, ring in comb(1) if ring == 1])
    step(s, [hexagon(s, *cell_xy(cx, cy, R, q, r), R * 0.92, BLUE) for q, r in sorted(blue)])


def s02_honeybee():
    s = slide(2, "THE HONEYBEE")
    title(s, "A honeybee attacked at a flower warns the hive off that path.", M, 1.45, 6.2, 2.5, size=40)
    text(s, M, 4.35, 6.4, 0.6, "One bee pays. The hive doesn't.", size=28, color=HONEY, track=-0.03)

    cx, cy, R = 8.95, 4.95, 0.36
    origin = (1, -1)
    for q, r, _ in comb(2):
        hexagon(s, *cell_xy(cx, cy, R, q, r), R * 0.92, CELL, CELL_LINE)
    text(s, 7.3, 6.5, 2, 0.18, "THE HIVE", size=8.5, color=MUTE, font=MONO, track=0.16)

    # The forager's path out to the flower, as a dashed curve.
    ox, oy = cell_xy(cx, cy, R, *origin)
    p0, p1, p2 = (ox + 0.1, oy - 0.3), (9.75, 2.35), (11.55, 2.12)
    pts = []
    for i in range(17):
        t = i / 16
        pts.append((
            (1 - t) ** 2 * p0[0] + 2 * t * (1 - t) * p1[0] + t * t * p2[0],
            (1 - t) ** 2 * p0[1] + 2 * t * (1 - t) * p1[1] + t * t * p2[1],
        ))
    fx, fy = 12.0, 1.98
    step(s, [
        polyline(s, pts, MUTE, 1.0, dash=True),
        dot(s, pts[9][0], pts[9][1], 0.055, WHITE),
        hexagon(s, fx, fy, 0.42, None, LINE2, 1.0),
        text(s, 9.2, 2.3, 2.2, 0.18, "FORAGER'S PATH", size=8.5, color=MUTE, font=MONO, track=0.16),
        text(s, 10.4, 1.38, 1.1, 0.18, "FLOWER", size=8.5, color=MUTE, font=MONO, track=0.16, align="r"),
    ])
    step(s, [
        hexagon(s, fx, fy, 0.25, RED),
        text(s, 10.9, 2.52, 1.68, 0.18, "ATTACKED", size=8.5, color=RED, font=MONO, track=0.16, align="r"),
    ])
    # The stop signal spreads across the comb, ring by ring, from the bee that paid.
    def dist(q, r):
        dq, dr = q - origin[0], r - origin[1]
        return max(abs(dq), abs(dr), abs(dq + dr))
    step(s, [
        hexagon(s, ox, oy, R * 0.92, HONEY),
        text(s, 7.3, 3.2, 2.4, 0.18, "STOP SIGNAL", size=8.5, color=HONEY, font=MONO, track=0.16),
    ])
    step(s, [hexagon(s, *cell_xy(cx, cy, R, q, r), R * 0.92, "B87A1B") for q, r, _ in comb(2) if dist(q, r) == 1])
    step(s, [hexagon(s, *cell_xy(cx, cy, R, q, r), R * 0.92, "4A310B", "8A5C14") for q, r, _ in comb(2) if dist(q, r) == 2])


def s03_problem():
    s = slide(3, "THE PROBLEM")
    title(s, "Agents have no stop signal.", M, 1.3, CW, 0.95, size=56)
    walls = [
        ("STRIPE WEBHOOKS", "No signatures found matching the expected signature for payload"),
        ("SUPABASE", "new row violates row-level security policy"),
        ("NEXT.JS", "params should be awaited before using its properties"),
    ]
    for i, (vendor, err) in enumerate(walls):
        y0 = 2.75 + i * 0.78
        rule(s, M, y0, 8.0, y0)
        rect(s, M, y0 + 0.45, 0.09, 0.09, RED, None)
        text(s, M + 0.3, y0 + 0.14, 6.9, 0.18, vendor, size=8.5, color=MUTE, font=MONO, track=0.16)
        text(s, M + 0.3, y0 + 0.37, 6.95, 0.26, err, size=12, color=WHITE, font=MONO)
    rule(s, M, 5.09, 8.0, 5.09)
    text(s, M, 5.5, CW, 0.6, [[("Every agent pays again. ", {}), ("The vendor never finds out.", {"color": HONEY})]],
         size=28, color=WHITE, track=-0.03)

    # Agent after agent flying into the same red cell.
    hx, hy = 11.85, 3.92
    hexagon(s, hx, hy, 0.56, RED)
    text(s, 10.85, 4.62, 1.73, 0.18, "SAME CRASH SITE", size=8.5, color=RED, font=MONO, track=0.12, align="r")
    text(s, 8.6, 2.75, 2, 0.18, "AGENTS", size=8.5, color=MUTE, font=MONO, track=0.16)
    shades = ["3A3A3A", "5A5A5A", "8A8A8A", "C4C4C4", "FFFFFF"]
    for lane, y_start in enumerate([3.1, 3.5, 3.92, 4.34, 4.74]):
        dots = []
        for k in range(5):
            t = (k + (lane % 2) * 0.5) / 5.6
            dots.append(dot(s, 8.7 + t * (hx - 0.75 - 8.7), y_start + t * (hy - y_start), 0.035 + 0.007 * k, shades[k]))
        step(s, dots)


def s04_loop():
    s = slide(4, "THE LOOP")
    title(s, "The loop.", M, 1.3, CW, 0.95, size=56)
    cw, gap = 2.6, (CW - 4 * 2.6) / 3
    steps = [
        (RED, [("Agent goes down → ", {}), ("mayday", {"color": RED})]),
        (WHITE, [("Postgres finds the crash site", {})]),
        (HONEY, [("Briefing: the fix that worked, ", {}), ("official fix first", {"color": HONEY})]),
        (BLUE, [("Rescue", {"color": BLUE}), (" confirmed, the best fix rises", {})]),
    ]
    for i, (color, runs) in enumerate(steps):
        x = M + i * (cw + gap)
        parts = []
        if i:
            parts.append(rule(s, x - gap + 0.07, 4.25, x - 0.07, 4.25, WHITE, 1.0, arrow=True))
        parts += [
            rect(s, x, 2.55, cw, 3.4, None, LINE),
            rule(s, x, 2.55, x + cw, 2.55, color, 2.0),
            text(s, x + 0.3, 2.8, 1.0, 0.8, str(i + 1), size=50, color=WHITE, track=-0.04, line=0.9),
            hexagon(s, x + cw - 0.52, 3.17, 0.24, color if color != WHITE else None, color, 1.0),
            text(s, x + 0.3, 4.0, cw - 0.6, 1.8, [runs], size=20, color=WHITE, track=-0.015, line=1.0),
        ]
        step(s, parts)


def s05_live():
    s = slide(5, "LIVE")
    stats = ["TOTAL MAYDAYS", "RESCUES", "CRASH SITES", "VENDORS"]
    note = "includes charted sites; live and test-flight traffic is marked"
    get = f"GET {SITE}/api/v1/map"
    if RADAR.exists():
        # Screenshot on the right inside a hairline frame, the words on the left.
        title(s, "Watch one go down.", M, 1.3, 5.6, 1.6, size=48)
        text(s, M, 3.1, 5.6, 0.3, SITE, size=15, color=HONEY, font=MONO)
        for i, label in enumerate(stats):
            x, y = M + (i % 2) * 2.85, 3.75 + (i // 2) * 0.95
            rule(s, x, y, x + 2.6, y)
            text(s, x, y + 0.14, 2.6, 0.18, label, size=8.5, color=MUTE, font=MONO, track=0.16)
            text(s, x, y + 0.4, 2.6, 0.4, "live on screen", size=18, color=WHITE, track=-0.02)
        text(s, M, 5.85, 5.6, 0.18, get, size=8, color=MUTE, font=MONO)
        text(s, M, 6.15, 5.6, 0.3, note, size=11, color=MUTE)
        bx, by, bw, bh = 6.75, 1.3, W - M - 6.75, 5.3
        pic = s.shapes.add_picture(str(RADAR), 0, 0)
        ratio = pic.width / pic.height
        pw, ph = (bw - 0.16, (bw - 0.16) / ratio) if (bw - 0.16) / ratio <= bh - 0.16 else ((bh - 0.16) * ratio, bh - 0.16)
        px, py = bx + (bw - pw) / 2, by + (bh - ph) / 2
        pic.left, pic.top, pic.width, pic.height = Inches(px), Inches(py), Inches(pw), Inches(ph)
        rect(s, px - 0.08, py - 0.08, pw + 0.16, ph + 0.16, None, LINE2)
        return
    title(s, "Watch one go down.", M, 1.3, CW, 0.95, size=56)
    rect(s, M, 2.6, CW, 1.5, None, LINE2)
    hexagon(s, M + 0.47, 2.92, 0.09, RED)
    text(s, M + 0.68, 2.84, 4, 0.18, "LIVE NOW", size=8.5, color=MUTE, font=MONO, track=0.16)
    text(s, M + 0.35, 3.2, CW - 0.7, 0.65, SITE, size=30, color=HONEY, font=MONO, track=-0.01)
    cw = CW / 4
    for i, label in enumerate(stats):
        x = M + i * cw
        rule(s, x, 4.5, x + cw - 0.3, 4.5)
        text(s, x, 4.66, cw - 0.3, 0.18, label, size=8.5, color=MUTE, font=MONO, track=0.16)
        text(s, x, 4.95, cw - 0.3, 0.5, "live on screen", size=24, color=WHITE, track=-0.025)
    text(s, M, 5.85, CW, 0.2, get, size=10, color=MUTE, font=MONO)
    text(s, M, 6.2, CW, 0.3, note, size=13, color=MUTE)


def s06_vendor():
    s = slide(6, "THE VENDOR SIDE")
    title(s, "Every vendor gets a tower.", M, 1.3, CW, 0.95, size=54)
    points = [
        [("See where agents crash on your product, ranked by agents down and hours lost", {})],
        [("Pin the official fix at the exact crash site", {})],
        [("Pay per rescue, through Stripe. ", {}), ("Only when the fix works.", {"color": HONEY})],
    ]
    for i, runs in enumerate(points):
        y0 = 2.6 + i * 1.3
        rule(s, M, y0, 5.9, y0)
        text(s, M, y0 + 0.17, 1, 0.18, f"{i + 1:02d}", size=9, color=HONEY, font=MONO, track=0.16)
        text(s, M + 0.6, y0 + 0.14, 4.55, 0.95, [runs], size=18, color=WHITE, track=-0.015, line=1.0)
    rule(s, M, 6.5, 5.9, 6.5)

    # The tower, sketched: crash sites ranked, bars for agents down and hours lost.
    tx, ty, tw = 6.5, 2.6, W - M - 6.5
    rect(s, tx, ty, tw, 3.9, None, LINE2)
    for label, off in [("#", 0.25), ("CRASH SITE", 0.75), ("AGENTS DOWN", 3.75), ("HOURS LOST", 4.95)]:
        text(s, tx + off, ty + 0.15, 1.3, 0.16, label, size=7.5, color=MUTE, font=MONO, track=0.14)
    names, down, lost = [1.35, 2.1, 1.8, 2.3, 1.6], [1.0, 0.8, 0.58, 0.4, 0.24], [0.9, 0.62, 0.5, 0.3, 0.2]
    for i in range(5):
        y0 = ty + 0.45 + i * 0.69
        rule(s, tx, y0, tx + tw, y0)
        text(s, tx + 0.25, y0 + 0.25, 0.4, 0.18, f"{i + 1:02d}", size=9, color=MUTE, font=MONO)
        rect(s, tx + 0.75, y0 + 0.3, names[i], 0.09, "3A3A3A", None)
        rect(s, tx + 3.75, y0 + 0.3, down[i], 0.09, RED, None)
        rect(s, tx + 4.95, y0 + 0.3, lost[i], 0.09, MUTE, None)
    y0 = ty + 0.45
    step(s, [
        rect(s, tx, y0, 0.05, 0.69, HONEY, None),
        rect(s, tx + 2.3, y0 + 0.21, 1.25, 0.27, BLACK, HONEY),
        hexagon(s, tx + 2.45, y0 + 0.345, 0.065, HONEY),
        text(s, tx + 2.58, y0 + 0.285, 0.95, 0.14, "OFFICIAL FIX", size=7.5, color=HONEY, font=MONO, track=0.08),
    ])


def s07_airworthiness():
    s = slide(7, "AIRWORTHINESS")
    title(s, "A rating that can't be bought.", M, 1.3, CW, 0.9, size=50)
    text(s, M, 2.35, 9.5, 0.8, "Computed from real crashes and rescues. It moves only when agents stop going down.",
         size=20, color=MUTE, track=-0.015, line=1.0)
    grades = [("C", MUTE, 5), ("B", WHITE, 9), ("A", HONEY, 14)]
    R = 0.2
    for i, (letter, color, filled) in enumerate(grades):
        bx = M + i * 4.05
        parts = []
        if i:
            parts.append(rule(s, bx - 0.62, 4.1, bx - 0.2, 4.1, WHITE, 1.0, arrow=True))
        parts.append(text(s, bx, 3.4, 1.2, 1.4, letter, size=96, color=color, track=-0.04, line=0.85))
        for k in range(15):
            row, col = divmod(k, 5)
            x = bx + 1.45 + col * R * SQ3 + (row % 2) * R * SQ3 / 2
            y = 3.8 + row * R * 1.5
            on = k < filled
            parts.append(hexagon(s, x, y, R * 0.9, HONEY if on else CELL, None if on else CELL_LINE))
        step(s, parts)
    step(s, [
        rect(s, M, 5.62, 3.9, 0.55, None, LINE2),
        hexagon(s, M + 0.3, 5.895, 0.11, HONEY),
        text(s, M + 0.55, 5.8, 3.25, 0.22, [[("MAYDAY AIRWORTHINESS · ", {}), ("B 62", {"color": HONEY})]],
             size=11, color=WHITE, font=MONO),
    ])
    text(s, M + 4.3, 5.73, CW - 4.3, 0.4,
         [[("Agents check it before they build: ", {}), ("mayday_preflight", {"font": MONO, "size": 16, "color": HONEY, "track": 0}), (".", {})]],
         size=18, color=WHITE, track=-0.015)


def s08_built():
    s = slide(8, "HOW IT'S BUILT")
    title(s, "How it's built.", M, 1.3, CW, 0.95, size=54)
    cols = [
        ("SUPABASE", ["Postgres does the matching (pg_trgm + error codes).", "RLS is the permission model.", "Realtime is the UI."]),
        ("VERCEL", ["Next.js 16, the MCP server, deploys in seconds."]),
        ("STRIPE", ["Checkout to claim an airspace.", "Billing Meters for pay-per-rescue."]),
        ("CLAUDE", ["Claude Code plugin hook, six MCP tools, test flights flown by real agents."]),
    ]
    gap = 0.4
    cw = (CW - 3 * gap) / 4
    for i, (name, body) in enumerate(cols):
        x = M + i * (cw + gap)
        step(s, [
            rule(s, x, 2.7, x + cw, 2.7, LINE2),
            hexagon(s, x + 0.1, 3.0, 0.1, HONEY),
            text(s, x + 0.32, 2.905, cw - 0.32, 0.22, name, size=11.5, color=WHITE, font=MONO, track=0.16),
            text(s, x, 3.45, cw, 2.2, body, size=16, color=INK, track=-0.01, line=1.02, gap=7),
        ])
    rule(s, M, 5.9, W - M, 5.9)
    hexagon(s, M + 0.08, 6.25, 0.08, HONEY)
    text(s, M + 0.3, 6.13, CW - 0.3, 0.3, "A mayday is one SQL transaction.", size=13, color=HONEY, font=MONO)


def s09_business():
    s = slide(9, "THE BUSINESS")
    title(s, "Vendors already pay to stop developers failing on their product.", M, 1.3, CW, 1.55, size=44)
    rows = [
        ("Find", "test flights + live maydays"),
        ("Fix", "the official fix, delivered at the moment of failure"),
        ("Prove", "airworthiness and a per-rescue bill"),
    ]
    for i, (verb, rest) in enumerate(rows):
        y0 = 3.15 + i * 0.8
        step(s, [
            rule(s, M, y0, W - M, y0),
            text(s, M, y0 + 0.17, 1.7, 0.5, verb, size=26, color=HONEY, track=-0.03),
            text(s, M + 1.9, y0 + 0.2, CW - 1.9, 0.45, rest, size=22, color=WHITE, track=-0.02),
        ])
    rule(s, M, 5.55, W - M, 5.55)
    text(s, M, 5.9, CW, 0.5, [[("Day-one value with zero network: ", {}), ("launch a test flight.", {"color": HONEY})]],
         size=24, color=WHITE, track=-0.025)


def s10_close():
    s = slide(10, "CLOSE")
    title(s, "Every agent that goes down should be the last one to go down there.", M, 1.3, 11.4, 2.25, size=48)
    text(s, M, 3.85, CW, 0.35, SITE, size=18, color=HONEY, font=MONO)
    text(s, M, 4.3, CW, 0.35, REPO, size=18, color=WHITE, font=MONO)

    # The whole comb starts red, then turns honey and blue from the right.
    R, rows, cols = 0.27, 3, 24
    dx = R * SQ3
    cells = [(row, col, M + dx / 2 + col * dx + (row % 2) * dx / 2, 5.3 + row * R * 1.5) for row in range(rows) for col in range(cols)]
    for _, _, x, y in cells:
        hexagon(s, x, y, R * 0.9, RED)
    def tone(row, col):
        f = col + ((row * 7 + col * 13) % 5 - 2) * 0.6
        return None if f < 3.2 else HONEY if f < 7.4 else BLUE
    for band in range(5):
        lo, hi = band * 5, band * 5 + 5
        group = [hexagon(s, x, y, R * 0.9, tone(row, col)) for row, col, x, y in cells if lo <= col < hi and tone(row, col)]
        step(s, group)


# --------------------------------------------------------- motion (raw XML)

def add_motion(s):
    """Fade transition on every slide, plus the fade-in steps that play on their
    own once the slide is up. The XML mirrors what PowerPoint writes itself."""
    sld = s._element
    anchor = sld.find(qn("p:clrMapOvr"))
    transition = etree.fromstring(f'<p:transition xmlns:p="{P_NS}" spd="med"><p:fade/></p:transition>')
    anchor.addnext(transition)  # schema order: clrMapOvr, transition, timing
    ids = STEPS.get(s.slide_id, [])
    if not ANIMATE or not ids:
        return
    n, t, pars = 3, 0, []
    wait, dur = 350, 500
    for spid in ids:
        pars.append(
            f'<p:par><p:cTn id="{n + 1}" fill="hold"><p:stCondLst><p:cond delay="{t}"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{n + 2}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="afterEffect">'
            f'<p:stCondLst><p:cond delay="{wait}"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{n + 3}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
            f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{n + 4}" dur="{dur}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
        )
        n += 4
        t += wait + dur
    timing = etree.fromstring(
        f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        f'<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        f'{"".join(pars)}'
        f'</p:childTnLst></p:cTn></p:par>'
        f'</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    )
    transition.addnext(timing)


# -------------------------------------------------------------------- verify

def walk(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type is not None and hasattr(shp, "shapes"):
            yield from walk(shp.shapes)


def verify(path):
    deck = Presentation(str(path))
    assert len(deck.slides) == 10, f"expected 10 slides, got {len(deck.slides)}"
    sw, sh = deck.slide_width, deck.slide_height
    problems = []
    for i, s in enumerate(deck.slides, 1):
        head = s.shapes.title.text_frame.text
        notes = s.notes_slide.notes_text_frame.text
        kids = [etree.QName(c).localname for c in s._element]
        assert "transition" in kids, f"slide {i}: no transition"
        assert kids.index("transition") == kids.index("clrMapOvr") + 1, f"slide {i}: transition misplaced"
        if "timing" in kids:
            assert kids.index("timing") == kids.index("transition") + 1, f"slide {i}: timing misplaced"
        assert notes.strip(), f"slide {i}: no speaker notes"
        boxes = 0
        for shp in walk(s.shapes):
            if shp.left < 0 or shp.top < 0 or shp.left + shp.width > sw or shp.top + shp.height > sh:
                problems.append(f"slide {i}: '{shp.name}' leaves the slide")
            if shp.has_text_frame and shp.text_frame.text.strip():
                boxes += 1
        steps = len(s._element.findall(".//" + qn("p:animEffect")))
        print(f"{i:>2}  {head[:58]:<58}  notes {len(notes):>3} chars  text boxes {boxes:>2}  fade-ins {steps}")
    words = sum(len(s.notes_slide.notes_text_frame.text.split()) for s in deck.slides)
    print(f"speaker notes: {words} words, about {words / 150:.1f} minutes spoken")
    for p in problems + WARNINGS:
        print("WARNING:", p)
    assert not problems, "shapes extend beyond the slide"
    print(f"ok: {path.relative_to(ROOT)} ({path.stat().st_size // 1024} KB), 10 slides, all text inside the slide")


def main():
    master_bg = prs.slide_master.background.fill
    master_bg.solid()
    master_bg.fore_color.rgb = rgb(BLACK)
    for build in (s01_title, s02_honeybee, s03_problem, s04_loop, s05_live, s06_vendor, s07_airworthiness, s08_built, s09_business, s10_close):
        build()
    for s in prs.slides:
        add_motion(s)
    prs.core_properties.title = "Mayday: the stop signal for agents"
    prs.core_properties.author = "Mayday"
    prs.core_properties.subject = "Supabase Select 2026 hackathon"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    verify(OUT)


if __name__ == "__main__":
    main()
